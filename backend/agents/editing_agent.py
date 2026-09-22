import os
from pathlib import Path
import time

from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from pptx import Presentation

from agents.schemas import EditingResult
from agents.version_manager import save_version


load_dotenv()

model = ChatOpenAI(
    model="nvidia/nemotron-3-super-120b-a12b",
    api_key=os.getenv("NVIDIA_API_KEY"),
    base_url="https://integrate.api.nvidia.com/v1",
    temperature=0
)

structured_model = model.with_structured_output(EditingResult)


OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)


def generate_slide_content(instruction: str) -> EditingResult:

    prompt = f"""
You are a PowerPoint editing agent.

The user wants to modify an existing presentation.

USER REQUEST:
{instruction}

Generate the content for the new slides.

Rules:
1. Create exactly the number of slides requested by the user.
2. Each slide must have a clear title.
3. Each slide should contain 3-5 concise bullet points.
4. Do not invent unnecessary information.
5. Keep the content professional and presentation-friendly.

Return structured slide content.
"""

    return structured_model.invoke(prompt)


def edit_presentation(
    input_path: str,
    instruction: str,
    output_filename: str | None = None
):

    presentation = Presentation(input_path)

    # Ask the LLM for slide content
    editing_result = generate_slide_content(instruction)

    layout = presentation.slide_layouts[1]

    # Add generated slides
    for generated_slide in editing_result.slides:

        slide = presentation.slides.add_slide(layout)

        slide.shapes.title.text = generated_slide.title

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        for index, bullet in enumerate(generated_slide.bullets):

            paragraph = (
                text_frame.paragraphs[0]
                if index == 0
                else text_frame.add_paragraph()
            )

            paragraph.text = bullet
            paragraph.level = 0

    if output_filename is None:
        timestamp = int(time.time())
        output_filename = f"edited_presentation_v{timestamp}.pptx"

    output_path = OUTPUT_DIR / output_filename

    presentation.save(output_path)
    save_version(str(output_path))

    return str(output_path)