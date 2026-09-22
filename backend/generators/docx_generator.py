from pathlib import Path
from docx import Document
from docx.shared import Pt


OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)


def generate_docx(research_result, filename="research_report.docx"):

    document = Document()

    # -----------------------------
    # Title
    # -----------------------------

    title = document.add_heading(
        research_result.topic,
        level=0
    )

    # -----------------------------
    # Summary
    # -----------------------------

    document.add_heading("Summary", level=1)

    document.add_paragraph(
        research_result.summary
    )

    # -----------------------------
    # Key Findings
    # -----------------------------

    document.add_heading("Key Findings", level=1)

    for finding in research_result.key_findings:

        document.add_paragraph(
            finding,
            style="List Bullet"
        )

    # -----------------------------
    # Sources
    # -----------------------------

    document.add_heading("Sources", level=1)

    for source in research_result.sources:

        paragraph = document.add_paragraph()

        paragraph.add_run(
            source.title
        ).bold = True

        paragraph.add_run(
            f"\n{source.url}"
        )

        paragraph.add_run(
            f"\n{source.relevance}"
        )

    # -----------------------------
    # Save
    # -----------------------------

    output_path = OUTPUT_DIR / filename

    document.save(output_path)

    return str(output_path)