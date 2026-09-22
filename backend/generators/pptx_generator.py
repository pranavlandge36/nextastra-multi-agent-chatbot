from pathlib import Path
from pptx import Presentation


OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)


def generate_pptx(research_result, filename="research_presentation.pptx"):

    presentation = Presentation()

    # -----------------------------
    # Title slide
    # -----------------------------

    title_slide_layout = presentation.slide_layouts[0]

    slide = presentation.slides.add_slide(title_slide_layout)

    slide.shapes.title.text = research_result.topic

    if slide.placeholders and len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Research Report"

    # -----------------------------
    # Summary slide
    # -----------------------------

    content_layout = presentation.slide_layouts[1]

    slide = presentation.slides.add_slide(content_layout)

    slide.shapes.title.text = "Summary"

    text_frame = slide.placeholders[1].text_frame
    text_frame.text = research_result.summary

    # -----------------------------
    # Key findings
    # -----------------------------

    slide = presentation.slides.add_slide(content_layout)

    slide.shapes.title.text = "Key Findings"

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    for index, finding in enumerate(
        research_result.key_findings
    ):

        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.text = finding
        paragraph.level = 0

    # -----------------------------
    # Sources
    # -----------------------------

    slide = presentation.slides.add_slide(content_layout)

    slide.shapes.title.text = "Sources"

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    for index, source in enumerate(
        research_result.sources
    ):

        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.text = (
            f"{source.title}\n"
            f"{source.url}"
        )

    # -----------------------------
    # Save
    # -----------------------------

    output_path = OUTPUT_DIR / filename

    presentation.save(output_path)

    return str(output_path)