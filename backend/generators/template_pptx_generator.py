from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)


def generate_pptx_from_template(
    research_result,
    template_path: str,
    filename: str = "generated_presentation.pptx"
):

    # Load template
    presentation = Presentation(template_path)

    topic = research_result.topic
    summary = research_result.summary
    findings = research_result.key_findings
    sources = research_result.sources

    # --------------------------------------------------
    # Remove old images from template
    # --------------------------------------------------

    for slide in presentation.slides:

        for shape in list(slide.shapes):

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                sp = shape._element
                sp.getparent().remove(sp)

    # --------------------------------------------------
    # Create content for all template slides
    # --------------------------------------------------

    slide_contents = []

    # Slide 1
    slide_contents.append({
        "title": topic,
        "body": summary
    })

    # Slides 2 onwards
    for i, finding in enumerate(findings):

        slide_contents.append({
            "title": f"{topic} — Key Point {i + 1}",
            "body": finding
        })

    # Sources slide
    source_text = "\n".join(
        f"• {source.title}\n  {source.url}"
        for source in sources
    )

    slide_contents.append({
        "title": "Sources",
        "body": source_text
    })

    # Conclusion slide
    slide_contents.append({
        "title": "Conclusion",
        "body": summary
    })

    # --------------------------------------------------
    # Fill existing template slides
    # --------------------------------------------------

    for index, slide in enumerate(presentation.slides):

        if index >= len(slide_contents):
            break

        content = slide_contents[index]

        update_slide_text(
            slide,
            content["title"],
            content["body"]
        )

    # --------------------------------------------------
    # Save
    # --------------------------------------------------

    output_path = OUTPUT_DIR / filename

    presentation.save(output_path)

    return str(output_path)


def clear_text_from_shape(shape):
    # Normal text shape
    if shape.has_text_frame:
        shape.text_frame.clear()

    # Text inside grouped shapes
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            clear_text_from_shape(child)


def update_slide_text(slide, title, body):

    # Remove all existing text from the slide
    for shape in slide.shapes:
        clear_text_from_shape(shape)

    # Find a suitable shape for title
    title_shape = None

    for shape in slide.shapes:
        if shape.has_text_frame and shape.shape_type != MSO_SHAPE_TYPE.GROUP:
            title_shape = shape
            break

    # If no text shape exists, create one
    if title_shape is None:
        title_shape = slide.shapes.add_textbox(
            600000,
            500000,
            8500000,
            1000000
        )

    title_shape.text_frame.text = title

    # Create a fresh body textbox
    body_box = slide.shapes.add_textbox(
        700000,
        1800000,
        8500000,
        4000000
    )

    body_box.text_frame.text = body