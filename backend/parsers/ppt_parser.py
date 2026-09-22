from pptx import Presentation


def extract_pptx_content(file_path: str):

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_data = {
            "slide_number": slide_number,
            "layout": slide.slide_layout.name,
            "texts": [],
            "fonts": [],
            "images": 0
        }

        for shape in slide.shapes:

            # Extract text
            if hasattr(shape, "text") and shape.text.strip():

                slide_data["texts"].append(
                    shape.text.strip()
                )

                # Extract font information
                if shape.has_text_frame:

                    for paragraph in shape.text_frame.paragraphs:

                        for run in paragraph.runs:

                            if run.font.name:
                                slide_data["fonts"].append(
                                    run.font.name
                                )

            # Count images
            if shape.shape_type == 13:
                slide_data["images"] += 1

        slides.append(slide_data)

    return {
        "file_type": "pptx",
        "slide_count": len(presentation.slides),
        "slides": slides
    }